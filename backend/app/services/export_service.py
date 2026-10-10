import io
from datetime import date
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import Document
from app.services import cross_checks
from app.services.candidates import anonymize, search

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD_FILL = PatternFill("solid", start_color="2147E0")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
FAMILIES = {
    "software": "Développement logiciel", "data_ai": "Data & IA", "it_ops": "Infrastructure & DevOps",
    "product_design": "Produit & Design", "finance_accounting": "Finance & Comptabilité", "hr": "Ressources humaines",
    "sales_marketing": "Commercial & Marketing", "operations": "Opérations & Logistique", "engineering": "Ingénierie & Industrie",
    "healthcare": "Santé", "legal": "Juridique", "education": "Éducation & Formation", "admin_support": "Administratif & Support",
    "other": "Autre",
}


def clean_data(d: dict) -> dict:
    return {k: v for k, v in d.items() if not k.startswith("_") and k != "derived"}


def _date(s):
    try:
        return date.fromisoformat(s)
    except (TypeError, ValueError):
        return s


def _sheet(wb: Workbook, title: str, headers: list, rows: List[list], widths: list, money=(), dates=()) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
        for c in ws[ws.max_row]:
            if isinstance(c.value, str) and c.value.startswith(FORMULA_STARTS):
                c.data_type = "s"  # le texte reste du texte : jamais interprété comme une formule
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in money:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = "#,##0.00"
    for col in dates:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def select_invoices(db: Session, company_id: int, date_from: Optional[date], date_to: Optional[date], validated_only: bool) -> List[Document]:
    docs = []
    for doc in cross_checks.invoices_of(db, company_id):
        d = doc.extracted_data
        if validated_only and not (d.get("_review") or {}).get("validated"):
            continue
        if date_from or date_to:
            inv = _date(d.get("invoice_date"))
            if not isinstance(inv, date) or (date_from and inv < date_from) or (date_to and inv > date_to):
                continue
        docs.append(doc)
    return sorted(docs, key=lambda x: (str(x.extracted_data.get("invoice_date") or ""), x.id))


def document_payload(doc: Document, anonymous: bool = False) -> dict:
    d = doc.extracted_data or {}
    rev, val = d.get("_review") or {}, d.get("_validation") or {}
    data = clean_data(d)
    if anonymous and doc.doc_type == "cv":
        data = anonymize.anonymize_cv(data, doc.id)
    return {
        "document_id": doc.id,
        "filename": None if anonymous else doc.original_filename,
        "type": doc.doc_type,
        "status": doc.status.value,
        "validated": bool(rev.get("validated")),
        "validated_at": rev.get("at"),
        "corrected_fields": rev.get("corrected_fields", []),
        "issues": val.get("issues", []),
        "data": data,
        "derived": d.get("derived"),
    }


def invoices_workbook(db: Session, company_id: int, date_from, date_to, validated_only: bool) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    invoices, lines, vat = [], [], []
    for doc in select_invoices(db, company_id, date_from, date_to, validated_only):
        d = doc.extracted_data
        val = d.get("_validation") or {}
        invoices.append([
            doc.id, doc.original_filename, "Avoir" if d.get("document_kind") == "credit_note" else "Facture",
            d.get("supplier_name"), d.get("supplier_siret"), d.get("supplier_vat_number"), d.get("customer_name"),
            d.get("invoice_number"), _date(d.get("invoice_date")), _date(d.get("due_date")), d.get("currency"),
            d.get("total_ht"), d.get("total_vat"), d.get("total_ttc"), d.get("iban"),
            "Oui" if (d.get("_review") or {}).get("validated") else "Non", " | ".join(val.get("issues") or []),
        ])
        for ln in d.get("lines") or []:
            lines.append([doc.id, d.get("invoice_number"), d.get("supplier_name"), ln.get("description"),
                          ln.get("quantity"), ln.get("unit_price"), ln.get("vat_rate"), ln.get("total_ht")])
        for v in d.get("vat_breakdown") or []:
            vat.append([doc.id, d.get("invoice_number"), d.get("supplier_name"), v.get("rate"), v.get("base"), v.get("amount")])

    _sheet(wb, "Factures",
           ["Document", "Fichier", "Type", "Fournisseur", "SIRET", "N° TVA", "Client", "N° facture", "Date", "Échéance",
            "Devise", "Total HT", "TVA", "Total TTC", "IBAN", "Validé", "Alertes"],
           invoices, [10, 28, 10, 28, 18, 18, 24, 18, 12, 12, 8, 13, 13, 13, 30, 8, 50], money=(12, 13, 14), dates=(9, 10))
    _sheet(wb, "Lignes", ["Document", "N° facture", "Fournisseur", "Désignation", "Quantité", "Prix unitaire HT", "TVA (%)", "Total HT"],
           lines, [10, 18, 28, 50, 10, 16, 10, 13], money=(5, 6, 8))
    _sheet(wb, "TVA", ["Document", "N° facture", "Fournisseur", "Taux (%)", "Base HT", "Montant TVA"],
           vat, [10, 18, 28, 10, 14, 14], money=(5, 6))
    return _bytes(wb)


def invoices_json(db: Session, company_id: int, date_from, date_to, validated_only: bool) -> list:
    return [document_payload(d) for d in select_invoices(db, company_id, date_from, date_to, validated_only)]


def candidates_workbook(db: Session, company_id: int, anonymous: bool) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for p, d in search.load_candidates(db, company_id):
        data = d.extracted_data or {}
        rows.append([
            d.id, anonymize.display_name(d.id, data, anonymous), data.get("headline"), data.get("location"),
            round(p.years or 0, 1), FAMILIES.get(p.job_family or "", p.job_family),
            ", ".join(p.skills or []),
            ", ".join(l.get("language", "") + (f" ({l['level']})" if l.get("level") else "") for l in data.get("languages") or []),
            None if anonymous else data.get("email"), None if anonymous else data.get("phone"),
        ])
    _sheet(wb, "Candidats",
           ["Document", "Candidat", "Titre", "Localisation", "Années d'expérience", "Famille de métier", "Compétences", "Langues", "E-mail", "Téléphone"],
           rows, [10, 26, 34, 18, 12, 24, 60, 28, 28, 18])
    return _bytes(wb)


def candidates_json(db: Session, company_id: int, anonymous: bool) -> list:
    return [document_payload(d, anonymous) for _, d in search.load_candidates(db, company_id)]