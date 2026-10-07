import io
import re

from docx import Document as Docx
from docx.shared import Pt, RGBColor

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GRAY = RGBColor(0x66, 0x6E, 0x85)


def _ym(v) -> str:
    m = re.fullmatch(r"(\d{4})(?:-(\d{2}))?", str(v or ""))
    if not m:
        return str(v or "")
    return f"{m.group(2)}/{m.group(1)}" if m.group(2) else m.group(1)


def _line(doc, text: str, size: int = 10, bold: bool = False, italic: bool = False, color=None):
    run = doc.add_paragraph().add_run(text)
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    if color:
        run.font.color.rgb = color


def build_docx(data: dict, name: str, company: str, years: float) -> bytes:
    doc = Docx()
    doc.sections[0].header.paragraphs[0].text = f"Dossier de compétences · {company}"

    _line(doc, name, 22, bold=True)
    if data.get("headline"):
        _line(doc, data["headline"], 13, color=GRAY)
    meta = " · ".join(x for x in (data.get("location"), f"{years:g} ans d'expérience" if years else None) if x)
    if meta:
        _line(doc, meta, 10, color=GRAY)

    if data.get("summary"):
        doc.add_heading("Profil", level=2)
        doc.add_paragraph(data["summary"])

    skill_names = [s["name"] for s in data.get("skills") or [] if s.get("name")]
    if skill_names:
        doc.add_heading("Compétences", level=2)
        doc.add_paragraph(" · ".join(skill_names))

    if data.get("experiences"):
        doc.add_heading("Expériences professionnelles", level=2)
        for e in data["experiences"]:
            _line(doc, " — ".join(x for x in (e.get("title"), e.get("company")) if x) or "Expérience", 11, bold=True)
            end = "aujourd'hui" if e.get("is_current") else _ym(e.get("end_date"))
            period = " → ".join(x for x in (_ym(e.get("start_date")), end) if x)
            when = " · ".join(x for x in (period, e.get("location")) if x)
            if when:
                _line(doc, when, 9, italic=True, color=GRAY)
            for row in re.split(r"\n+|\s*•\s*", e.get("description") or ""):
                if row.strip(" -–"):
                    doc.add_paragraph(row.strip(" -–"), style="List Bullet")

    if data.get("education"):
        doc.add_heading("Formations", level=2)
        for e in data["education"]:
            text = ", ".join(x for x in (e.get("degree"), e.get("field"), e.get("institution")) if x)
            doc.add_paragraph(f"{text} ({e['year']})" if e.get("year") else text, style="List Bullet")

    if data.get("languages"):
        doc.add_heading("Langues", level=2)
        doc.add_paragraph(" · ".join(l["language"] + (f" ({l['level']})" if l.get("level") else "") for l in data["languages"]))

    if data.get("certifications"):
        doc.add_heading("Certifications", level=2)
        for c in data["certifications"]:
            doc.add_paragraph(" — ".join(x for x in (c.get("name"), c.get("issuer"), str(c["year"]) if c.get("year") else None) if x), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()